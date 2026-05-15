"""
Bulk import parser — Excel/CSV savol fayllarni import qilish
"""

import csv
import io
from typing import Any

from openpyxl import load_workbook

from apps.catalog.models import ImportBatch, QuestionDraft, Subject, Topic
from apps.organizations.models import Organization
from core.tenant import tenant_context


class BulkImportParser:
    """Excel/CSV fayldan savol va variantlarni oqish."""

    def __init__(self, file, file_type: str, organization: Organization):
        self.file = file
        self.file_type = file_type
        self.organization = organization
        self.errors = []
        self.rows = []

    def parse_excel(self) -> list[dict[str, Any]]:
        """Excel faylni oqish."""
        try:
            workbook = load_workbook(self.file)
            worksheet = workbook.active

            rows = []
            headers = None

            for idx, row in enumerate(worksheet.iter_rows(values_only=True), 1):
                if idx == 1:
                    headers = row
                    continue

                row_data = dict(zip(headers, row, strict=False)) if headers else {}
                if any(row_data.values()):
                    rows.append(row_data)

            return rows
        except Exception as e:
            self.errors.append(f'Excel parse xatosi: {str(e)}')
            return []

    def parse_csv(self) -> list[dict[str, Any]]:
        """CSV faylni oqish."""
        try:
            content = self.file.read()
            if isinstance(content, bytes):
                content = content.decode('utf-8')

            reader = csv.DictReader(io.StringIO(content))
            return [row for row in reader if any(row.values())]
        except Exception as e:
            self.errors.append(f'CSV parse xatosi: {str(e)}')
            return []

    def parse(self) -> tuple[list[dict], list[str]]:
        """Fayl turini aniqlab parse qilish."""
        if self.file_type.lower() in ['xlsx', 'xls']:
            self.rows = self.parse_excel()
        elif self.file_type.lower() == 'csv':
            self.rows = self.parse_csv()
        else:
            self.errors.append(f"Noma'lum fayl turi: {self.file_type}")

        return self.rows, self.errors

    def _validate_row(self, row: dict, row_number: int) -> tuple[bool, list[str]]:
        """Bir qator savol ma'lumotlarini tekshirish."""
        errors = []

        # Zaruri maydonlar
        if not row.get('question_text') or not str(row.get('question_text')).strip():
            errors.append(f"Qator {row_number}: Savol matni bo'sh")

        if not row.get('type') or row.get('type') not in ['SC', 'MC', 'MT', 'OR', 'FB', 'OE', 'FU']:
            errors.append(f"Qator {row_number}: Noto'g'ri savol turi")

        if not row.get('subject') or not str(row.get('subject')).strip():
            errors.append(f'Qator {row_number}: Fan aniqlanmagan')

        # Variantlar
        if row.get('type') in ['SC', 'MC']:
            option_cols = [k for k in row.keys() if k.startswith('option_')]
            if not option_cols:
                errors.append(f"Qator {row_number}: Variantlar yo'q")

        return len(errors) == 0, errors

    def validate_and_create_drafts(self, batch: ImportBatch) -> tuple[int, list[str]]:
        """Savoklari draft sifatida qo'shish."""
        created_count = 0
        all_errors = []

        with tenant_context(self.organization):
            for row_number, row in enumerate(self.rows, 2):  # 2-chi qatordan boshlash (header skip)
                is_valid, errors = self._validate_row(row, row_number)

                # Subject va Topic topish
                subject = None
                topic = None
                try:
                    subject = Subject.objects.get(
                        name__iexact=row.get('subject', '').strip(), organization=self.organization
                    )
                    if row.get('topic'):
                        topic = Topic.objects.filter(
                            name__iexact=row.get('topic', '').strip(), subject=subject
                        ).first()
                except Subject.DoesNotExist:
                    errors.append(f"Qator {row_number}: Fan '{row.get('subject')}' topilmadi")
                    is_valid = False

                # Draft yaratish
                draft_data = {
                    'text': row.get('question_text', ''),
                    'explanation': row.get('explanation', ''),
                    'difficulty': row.get('difficulty', 'medium'),
                    'options': self._extract_options(row),
                }

                QuestionDraft.objects.create(
                    batch=batch,
                    data=draft_data,
                    type=row.get('type'),
                    subject=subject,
                    topic=topic,
                    is_valid=is_valid,
                    validation_errors=errors if errors else None,
                )

                if is_valid and not errors:
                    created_count += 1
                else:
                    all_errors.extend(errors)

        return created_count, all_errors

    def _extract_options(self, row: dict) -> list[dict]:
        """Variantlarni foydalanuvchi jadvaldagi kolonkalardan oqish."""
        options = []
        for key, value in row.items():
            if key.startswith('option_'):
                option_id = key.replace('option_', '')
                is_correct = row.get(f'correct_{option_id}', False)
                if value:
                    options.append(
                        {'id': option_id, 'text': str(value), 'is_correct': bool(is_correct)}
                    )
        return options


def bulk_import_from_file(
    file, file_type: str, organization: Organization, batch: ImportBatch
) -> dict[str, Any]:
    """Bulk import jarayoni."""
    parser = BulkImportParser(file, file_type, organization)
    rows, parse_errors = parser.parse()

    if parse_errors:
        batch.status = ImportBatch.Status.FAILED
        batch.error_log = parse_errors
        batch.save()
        return {'success': False, 'total': 0, 'processed': 0, 'errors': parse_errors}

    batch.total_questions = len(rows)
    batch.status = ImportBatch.Status.PROCESSING
    batch.save()

    created, errors = parser.validate_and_create_drafts(batch)

    batch.processed_questions = created
    batch.status = ImportBatch.Status.COMPLETED if not errors else ImportBatch.Status.COMPLETED
    batch.error_log = errors if errors else None
    batch.save()

    return {'success': True, 'total': len(rows), 'processed': created, 'errors': errors}
